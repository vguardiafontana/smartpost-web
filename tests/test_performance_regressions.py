from concurrent.futures import ThreadPoolExecutor
from copy import copy
import random

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side

from app.procesamiento.io_sources import rows_to_dicts
from app.procesamiento.rules import EXPECTED_HEADERS, normalize_key
from app.reportes.portable_report import _copy_cell_style


def original_mapping(headers, rows):
    # Contrato anterior: primera columna duplicada gana, filas vacías se omiten.
    columns = {}; duplicates = []
    for i, header in enumerate(headers):
        key = normalize_key(header)
        if not key: continue
        if key in columns: duplicates.append(header)
        else: columns[key] = i
    mapped = []
    for raw in rows:
        if all(v in (None, '') for v in raw): continue
        item = {}
        for expected in EXPECTED_HEADERS:
            index = columns.get(normalize_key(expected))
            item[expected] = raw[index] if index is not None and index < len(raw) else None
        mapped.append(item)
    return mapped, duplicates


def test_mapping_same_values_order_duplicates_and_irregular_rows():
    rng = random.Random(294)
    for _ in range(30):
        headers = rng.sample(EXPECTED_HEADERS, 25)+['Nro Guía', 'Remito', '', 'Estado', 'NRO GUIA']
        rows = [[rng.choice([None, '', '00123', '#ABC', 0, 1.25, 'a\nb', False]) for _ in range(rng.randrange(1,40))] for _ in range(20)]
        rows.extend([[], [None]*30, ['']*30])
        actual = rows_to_dicts(headers, rows)
        assert actual == original_mapping(headers, rows)
        assert list(actual[0][0]) == EXPECTED_HEADERS


def test_mapping_does_not_reuse_previous_files_or_mutate_rows():
    rows = [['1', 'second']]
    assert rows_to_dicts(['Nro Guia','Remito'], rows)[0][0]['Remito'] == 'second'
    assert rows_to_dicts(['Remito','Nro Guia'], rows)[0][0]['Remito'] == '1'
    assert rows == [['1', 'second']]


def test_style_copy_preserves_every_component_and_is_independent():
    wb = Workbook(); ws = wb.active
    source = ws['A1']; dest = ws['A2']
    source.font = Font(name='Arial', size=18, bold=True, italic=True, color='123456')
    source.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=False, text_rotation=45)
    source.fill = PatternFill('solid', fgColor='abcdef')
    source.border = Border(left=Side(style='thick', color='000000'))
    source.protection = Protection(locked=False, hidden=True)
    source.number_format = '000000'; source.quotePrefix = True
    _copy_cell_style(source,dest)
    import importlib.util
    from pathlib import Path
    spec = importlib.util.spec_from_file_location('reference_style', Path(__file__).parent/'reference/report_v293.py')
    reference = importlib.util.module_from_spec(spec); spec.loader.exec_module(reference)
    expected = ws['A3']
    reference._copy_cell_style(source, expected)
    assert dest._style == expected._style and source._style is not dest._style
    dest.number_format = '@'
    assert source.number_format == '000000'


def test_style_copy_across_workbooks_uses_destination_style_tables():
    source=Workbook().active['A1'];dest=Workbook().active['A1']
    source.font=Font(name='Example',bold=True)
    _copy_cell_style(source,dest)
    assert copy(dest.font) == copy(source.font)


def test_default_config_cache_reads_content_not_timestamp(tmp_path, monkeypatch):
    from app.core import config
    import yaml
    defaults=yaml.safe_load(config.DEFAULT_CONFIG_PATH.read_text(encoding='utf-8'))
    user=tmp_path/'user.yaml'; base=tmp_path/'defaults.yaml'
    user.write_text(yaml.safe_dump(defaults),encoding='utf-8')
    base.write_text(yaml.safe_dump(defaults),encoding='utf-8')
    monkeypatch.setattr(config,'ensure_user_config',lambda:user)
    monkeypatch.setattr(config,'DEFAULT_CONFIG_PATH',base)
    config._parse_defaults.cache_clear()
    config.load_config();config.load_config()
    assert config._parse_defaults.cache_info().hits >= 1
    # Cambio de defaults de igual longitud y mtime conservado.
    defaults['new_setting']='AAA';base.write_text(yaml.safe_dump(defaults),encoding='utf-8')
    assert config.load_config()['new_setting']=='AAA'
    import os
    previous=base.stat()
    defaults['new_setting']='BBB';base.write_text(yaml.safe_dump(defaults),encoding='utf-8')
    os.utime(base,ns=(previous.st_atime_ns,previous.st_mtime_ns))
    # Personalización del usuario prevalece, como antes.
    assert config.load_config()['new_setting']=='AAA'
    user_cfg=yaml.safe_load(user.read_text(encoding='utf-8'));user_cfg.pop('new_setting')
    user.write_text(yaml.safe_dump(user_cfg),encoding='utf-8')
    assert config.load_config()['new_setting']=='BBB'
    base.write_text('invalid: [',encoding='utf-8')
    with pytest.raises(yaml.YAMLError):config.load_config()


def test_cached_defaults_never_share_mutable_data_between_callers():
    from app.core.config import _parse_defaults, _deep_merge
    text='settings:\n  nested: [a, b]\n'
    def read_and_modify(i):
        value=_deep_merge(_parse_defaults(text),{})
        value['settings']['nested'].append(i)
        return value
    with ThreadPoolExecutor(max_workers=4) as pool:
        results=list(pool.map(read_and_modify,range(16)))
    assert all(len(r['settings']['nested'])==3 for r in results)
    assert _parse_defaults(text)=={'settings':{'nested':['a','b']}}


@pytest.mark.parametrize('size', [0, 1, 37, 210])
def test_xlsx_all_parts_equal_to_previous_engine(tmp_path, size):
    import importlib.util
    from pathlib import Path
    from datetime import datetime
    import zipfile
    import xml.etree.ElementTree as ET
    from app.reportes.portable_report import write_report_portable
    spec=importlib.util.spec_from_file_location('reference_report',Path(__file__).parent/'reference/report_v293.py')
    reference=importlib.util.module_from_spec(spec);spec.loader.exec_module(reference)
    rows=[]
    for i in range(size):
        row={h:None for h in EXPECTED_HEADERS}
        row.update({'Nro Guia':6200000+i,'Remito':['#ABC','000123',123,'2000018857041302'][i%4],
                    'Fecha Ingreso a Planta':datetime(2026,10,8,15,30),
                    'FECHA ULTIMO ESTADO':datetime(2026,10,8,20,30),'Fecha Planilla':datetime(2026,10,8),
                    'Estado':'1RVT - 1° VISITA','Servicio':'FLEX','Documento':'034.123','COM.CELULAR':'01112345',
                    'Domicilio':'Calle 123\nPiso 2','Trabajo':'000123'})
        rows.append(row)
    template=Path(__file__).resolve().parents[1]/'assets/TEMPLATE_REPORTE.xlsx'
    before=tmp_path/'before.xlsx';after=tmp_path/'after.xlsx'
    a=reference.write_report_portable(template,before,rows,'PRUEBA',{})
    b=write_report_portable(template,after,rows,'PRUEBA',{})
    assert a==b
    with zipfile.ZipFile(before) as az,zipfile.ZipFile(after) as bz:
        assert set(az.namelist())==set(bz.namelist())
        for name in az.namelist():
            left,right=az.read(name),bz.read(name)
            if name=='docProps/core.xml':
                def strip_time(value):
                    root=ET.fromstring(value)
                    for node in list(root):
                        if node.tag in ('{http://purl.org/dc/terms/}created','{http://purl.org/dc/terms/}modified'):root.remove(node)
                    return ET.tostring(root)
                left,right=strip_time(left),strip_time(right)
            assert left==right,name
