from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.procesamiento.rules import EXPECTED_HEADERS
from app.reportes.portable_report import write_report_portable
from app.visitas.portable_excel import convert_visits_csv_to_xlsx_portable
from app.core.models import RunContext


def _row():
    row = {h: None for h in EXPECTED_HEADERS}
    row.update({
        "Nro Guia": 6249087,
        "Fecha ingreso a Planta": datetime(2026, 10, 7, 18, 30),
        "Cliente": "ZETABURGOS SRL",
        "Estado": "RC - EN RUTA PARA SU ENTREGA",
        "Remito": "001234",
        "Servicio": "FLEX",
        "Destinatario": "Persona Prueba",
        "Domicilio": "Calle 123",
        "Localidad": "CABA",
        "Distribuidor Recorrido": "37424493APP - PRUEBA",
        "CP": 1000,
        "FECHA ULTIMO ESTADO": datetime(2026, 10, 7, 19, 0),
        "Zona Destino": "CABA C1",
    })
    return row


def test_portable_report_preserves_template_and_report_semantics(tmp_path):
    root = Path(__file__).resolve().parents[1]
    template = root / "assets" / "TEMPLATE_REPORTE.xlsx"
    out = tmp_path / "reporte.xlsx"
    audit = write_report_portable(template, out, [_row()], "ZETABURGOS", {})
    assert out.exists() and out.stat().st_size > 0
    wb = load_workbook(out, data_only=False)
    report = wb["REPORTE"]
    base = wb["base"]
    assert report["A1"].value == "Nro Guia"
    assert base["A2"].value == 6249087
    assert base["G2"].value == "001234"
    assert report["A2"].value.startswith("=IF(base!A2")
    assert report["F2"].value.startswith("=IF(D2")
    assert report.freeze_panes is None
    assert report.sheet_state == "visible"
    assert "ReporteTable" in report.tables
    assert report.tables["ReporteTable"].ref == "A1:Q2"
    assert len(report._images) >= 1
    assert report["A2"].font.name == "Calibri"
    assert report["A2"].font.sz == 11
    assert len(report.data_validations.dataValidation) == 1
    assert audit["excel"]["engine"] == "openpyxl-portable"


def test_portable_visitas_keeps_every_source_value_as_text(tmp_path, monkeypatch):
    csv = tmp_path / "visitas.csv"
    csv.write_text(";".join(EXPECTED_HEADERS) + "\n" + ";".join(["00123" if h == "Nro Guia" else "" for h in EXPECTED_HEADERS]) + "\n", encoding="utf-8")
    cfg = {"visitas": {"filename": "VISITAS {date_short}.xlsx"}, "paths": {"reports": "output/reportes"}}
    import app.visitas.portable_excel as mod
    monkeypatch.setattr(mod, "resolve_path", lambda _cfg, _key: tmp_path / "reports")
    run = RunContext("x", datetime(2026, 10, 7), tmp_path / "work")
    class L:
        def info(self, *a, **k): pass
    info = convert_visits_csv_to_xlsx_portable(cfg, run, csv, L())
    wb = load_workbook(info["path"], data_only=False)
    ws = wb["VISITAS"]
    assert ws["A2"].value == "00123"
    assert ws["A2"].number_format == "@"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref.endswith("2")

def test_web_reports_endpoints_and_download(monkeypatch):
    import threading, os
    os.environ.setdefault("SMARTPOST_WEB_PIN", "246810")
    os.environ.setdefault("SMARTPOST_WEB_SESSION_SECRET", "x" * 64)
    from fastapi.testclient import TestClient
    from app.web.security import WebSecuritySettings
    from app.web.server import create_app

    class FakeManager:
        def __init__(self): self.operation_lock=threading.Lock(); self.storage_backend="files"
        def start(self): pass
        def stop(self): pass
        def get_job(self, _): return None
        def create_job(self,*a,**k): raise AssertionError
        def retry_failed(self,*a,**k): raise AssertionError

    class FakeReports:
        storage_backend="files"
        def start(self): pass
        def stop(self): pass
        def create_job(self, date): return {"id":"r1","status":"queued","date":date}
        def get_job(self, jid): return {"id":jid,"status":"completed","artifacts":[{"id":"a1","filename":"A.xlsx","content_type":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}]}
        def retry(self, jid): return {"id":"r2","status":"queued","date":"2026-10-07"}
        def get_artifact(self, jid, aid): return ({"id":aid,"filename":"A.xlsx","content_type":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}, b"xlsx")

    app=create_app(security=WebSecuritySettings(pin="246810",session_secret="x"*64),manager=FakeManager(),report_manager=FakeReports(),preview_runner=lambda ids,sid:[],search_runner=lambda ids:[])
    with TestClient(app, base_url="https://testserver") as client:
        assert client.post('/api/auth/login',json={'pin':'246810'}).status_code==200
        r=client.post('/api/report-jobs',json={'date':'2026-10-07'})
        assert r.status_code==200 and r.json()['id']=='r1'
        d=client.get('/api/report-jobs/r1/artifacts/a1')
        assert d.status_code==200 and d.content==b'xlsx'
        assert 'attachment' in d.headers['content-disposition']


def test_web_home_has_reports_module():
    html=(Path(__file__).resolve().parents[1]/'app/web/static/index.html').read_text(encoding='utf-8')
    assert 'Reportes y Visitas' in html
    assert 'Copiar' in (Path(__file__).resolve().parents[1]/'app/web/static/app.js').read_text(encoding='utf-8')
