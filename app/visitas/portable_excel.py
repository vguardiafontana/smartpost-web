from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core.config import resolve_path
from app.core.models import RunContext
from app.core.utils import safe_filename
from app.procesamiento.io_sources import read_csv_file
from app.procesamiento.rules import EXPECTED_HEADERS, normalize_key


def _validated_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    headers, rows = read_csv_file(path)
    expected = [normalize_key(x) for x in EXPECTED_HEADERS]
    actual = [normalize_key(x) for x in headers]
    if len(headers) != len(EXPECTED_HEADERS):
        raise RuntimeError(f"VISITAS: esperaba {len(EXPECTED_HEADERS)} columnas y encontré {len(headers)}.")
    mismatches = [
        f"{i + 1}: '{EXPECTED_HEADERS[i]}' != '{headers[i]}'"
        for i in range(len(expected)) if expected[i] != actual[i]
    ]
    if mismatches:
        raise RuntimeError("VISITAS: encabezados inesperados: " + " | ".join(mismatches[:8]))
    clean_rows: list[list[str]] = []
    for row_num, row in enumerate(rows, start=2):
        if len(row) != len(headers):
            raise RuntimeError(f"VISITAS: fila {row_num} tiene {len(row)} columnas; esperaba {len(headers)}.")
        clean_rows.append(["" if value is None else str(value) for value in row])
    return [str(x) for x in headers], clean_rows


def _output_path(cfg: dict, run: RunContext) -> Path:
    visits_cfg = cfg.get("visitas") or {}
    template = str(visits_cfg.get("filename") or "VISITAS {date_short}.xlsx")
    name = template.format(date=run.message_date, date_short=run.date_short, date_ddmmyyyy=run.date_ddmmyyyy)
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    root = resolve_path(cfg, "reports") / run.date_ddmmyyyy
    root.mkdir(parents=True, exist_ok=True)
    return root / safe_filename(name)


def convert_visits_csv_to_xlsx_portable(cfg: dict, run: RunContext, csv_path: Path, logger) -> dict:
    headers, rows = _validated_csv(csv_path)
    output = _output_path(cfg, run)
    wb = Workbook()
    ws = wb.active
    ws.title = "VISITAS"
    for cidx, header in enumerate(headers, start=1):
        cell = ws.cell(1, cidx, header)
        cell.font = Font(bold=True)
        cell.number_format = "@"
    for ridx, row in enumerate(rows, start=2):
        for cidx, value in enumerate(row, start=1):
            cell = ws.cell(ridx, cidx, value)
            cell.number_format = "@"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    ws.freeze_panes = "A2"
    # Approximate Excel AutoFit without coercing values; cap exactly as desktop.
    for cidx in range(1, len(headers) + 1):
        max_len = max([len(str(ws.cell(r, cidx).value or "")) for r in range(1, len(rows) + 2)] or [0])
        ws.column_dimensions[get_column_letter(cidx)].width = min(45, max(8, max_len + 2))
    wb.save(output)
    with zipfile.ZipFile(output, "r") as zf:
        bad = zf.testzip()
        if bad:
            raise RuntimeError(f"VISITAS: XLSX corrupto ({bad}).")
    logger.info("[VISITAS] XLSX portable generado -> %s (%d registros)", output.name, len(rows))
    return {"status": "generated", "path": str(output), "rows": len(rows), "columns": len(headers)}
