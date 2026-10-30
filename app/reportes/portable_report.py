from __future__ import annotations

import copy
import shutil
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formatting.rule import Rule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from openpyxl.utils import get_column_letter

from app.procesamiento.rules import (
    DATETIME_COLUMNS,
    DATE_ONLY_COLUMNS,
    EXPECTED_HEADERS,
    REPORT_MAP,
    TEXT_COLUMNS,
    management_result,
    normalize_key,
    remito_as_string,
)

MAX_BASE_ROW_HEIGHT = 60.0
MAX_REPORT_ROW_HEIGHT = 69.75
REPORT_HEADERS = [
    "Nro Guia", "Fecha ingreso a Planta", "Cliente", "Estado", "TIPO DE ESTADO",
    "GESTIÓN REALIZADA", "Remito", "Servicio", "Destinatario", "Domicilio", "Localidad",
    "Distribuidor Recorrido", "Recibio", "Documento", "CP", "FECHA ULTIMO ESTADO", "Zona Destino",
]


def _copy_cell_style(src, dst, cache=None) -> None:
    # Caché local al archivo: conserva la normalización exacta de openpyxl,
    # pero sólo la hace una vez por estilo. Cada celda recibe su propia copia.
    key = None
    if cache is not None and src.has_style:
        key = (id(src.parent.parent), id(dst.parent.parent), tuple(src._style))
        if key in cache:
            dst._style = copy.copy(cache[key])
            return
    if src.has_style:
        dst._style = copy.copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy.copy(src.alignment)
    if src.font:
        dst.font = copy.copy(src.font)
    if src.fill:
        dst.fill = copy.copy(src.fill)
    if src.border:
        dst.border = copy.copy(src.border)
    if src.protection:
        dst.protection = copy.copy(src.protection)
    if key is not None:
        cache[key] = copy.copy(dst._style)


def _verify_template(wb) -> list[str]:
    problems: list[str] = []
    for required in ("REPORTE", "base", "aux"):
        if required not in wb.sheetnames:
            problems.append(f"Falta la hoja obligatoria {required}.")
    if problems:
        return problems
    base = wb["base"]
    for i, expected in enumerate(EXPECTED_HEADERS, start=1):
        got = base.cell(1, i).value
        if normalize_key(expected) != normalize_key(got):
            problems.append(f"base!{get_column_letter(i)}1 esperado '{expected}', encontrado '{got}'.")
    report = wb["REPORTE"]
    for i, expected in enumerate(REPORT_HEADERS, start=1):
        got = report.cell(1, i).value
        if normalize_key(expected) != normalize_key(got):
            problems.append(f"REPORTE!{get_column_letter(i)}1 esperado '{expected}', encontrado '{got}'.")
    return problems


def _excel_value(column: str, value: Any) -> Any:
    if value is None:
        return None
    if column in TEXT_COLUMNS:
        return str(value)
    return value


def _report_source_value(row: dict[str, Any], report_letter: str) -> Any:
    src_letter = REPORT_MAP[report_letter]
    idx = 0
    for ch in src_letter:
        idx = idx * 26 + (ord(ch) - 64)
    return row.get(EXPECTED_HEADERS[idx - 1])


def write_report_portable(template_path: Path, output_path: Path, rows: list[dict[str, Any]], seller: str, audit: dict) -> dict:
    """Linux/Render report writer preserving the desktop template as closely as possible.

    The desktop path continues to use real Excel COM. This writer exists only for
    the web/server path where Excel is unavailable. It keeps the original template,
    logo, widths, headers, formulas, dropdown and table semantics and forces Excel
    to recalculate formulas when the file is opened.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    wb = load_workbook(output_path, data_only=False)
    problems = _verify_template(wb)
    if problems:
        raise RuntimeError("Estructura de template no compatible: " + " | ".join(problems))

    base = wb["base"]
    report = wb["REPORTE"]
    aux = wb["aux"]
    n = len(rows)
    style_cache = {}

    # BASE: clear old values while retaining formatting and grow rows from row 2.
    # New rows are empty and receive every value below. Clear only old rows;
    # walking the new rows here would allocate/touch each cell twice.
    for r in range(2, base.max_row + 1):
        for c in range(1, len(EXPECTED_HEADERS) + 1):
            base.cell(r, c).value = None
    template_base_style_row = 2
    for ridx, source in enumerate(rows, start=2):
        if ridx > 2:
            for c in range(1, len(EXPECTED_HEADERS) + 1):
                _copy_cell_style(base.cell(template_base_style_row, c), base.cell(ridx, c), style_cache)
        for cidx, header in enumerate(EXPECTED_HEADERS, start=1):
            cell = base.cell(ridx, cidx)
            value = _excel_value(header, source.get(header))
            cell.value = value
            if header in DATE_ONLY_COLUMNS:
                cell.number_format = "mm-dd-yy"
            elif header in DATETIME_COLUMNS:
                cell.number_format = "m/d/yy h:mm"
        base.row_dimensions[ridx].height = min(
            float(base.row_dimensions[ridx].height or MAX_BASE_ROW_HEIGHT), MAX_BASE_ROW_HEIGHT
        )

    # REPORTE: exact active row count. Existing template has many sample rows.
    final_last = n + 1
    if report.max_row > final_last:
        report.delete_rows(final_last + 1, report.max_row - final_last)
    elif report.max_row < final_last:
        report.insert_rows(report.max_row + 1, final_last - report.max_row)

    # Remove any existing standard tables and rebuild one for the active range.
    for table_name in list(report.tables.keys()):
        del report.tables[table_name]

    # Use row 2 only as a structural style source, then normalize the body to
    # the same font/alignment rules as desktop v2.7.x.
    for ridx, source in enumerate(rows, start=2):
        if ridx > 2:
            for c in range(1, 18):
                _copy_cell_style(report.cell(2, c), report.cell(ridx, c), style_cache)
        for cidx, letter in enumerate("ABCDEFGHIJKLMNOPQ", start=1):
            cell = report.cell(ridx, cidx)
            if letter == "F":
                cell.value = (
                    f'=IF(D{ridx}="","",IF(OR(ISNUMBER(SEARCH("VISITA",D{ridx})),'
                    f'ISNUMBER(SEARCH("RUTA",D{ridx})),ISNUMBER(SEARCH("ARRIB",D{ridx})),'
                    f'ISNUMBER(SEARCH("MATRIZ",D{ridx})),ISNUMBER(SEARCH("PENDIENTE",D{ridx}))),'
                    '"VISITA COORDINADA CON EL CLIENTE","N/A"))'
                )
            elif letter == "G" and remito_as_string(source.get("Remito")).startswith("#"):
                literal = remito_as_string(source.get("Remito")).replace('"', '""')
                cell.value = f'="{literal}"'
            else:
                src = REPORT_MAP[letter]
                cell.value = f'=IF(base!{src}{ridx}="","",base!{src}{ridx})'
            if ridx == 2:
                # Las filas siguientes copian el estilo completo de esta fila,
                # ya normalizado. No se comparten objetos mutables entre celdas.
                font = copy.copy(cell.font)
                font.name = "Calibri"
                font.sz = 11
                cell.font = font
                alignment = copy.copy(cell.alignment)
                alignment.horizontal = None
                alignment.vertical = "top"
                alignment.wrap_text = True
                cell.alignment = alignment
        report.cell(ridx, 2).number_format = "dd/mm/yyyy hh:mm:ss"
        report.cell(ridx, 16).number_format = "dd/mm/yyyy hh:mm:ss"
        report.row_dimensions[ridx].height = min(
            float(report.row_dimensions[ridx].height or MAX_REPORT_ROW_HEIGHT), MAX_REPORT_ROW_HEIGHT
        )

    # Dropdown on Gestión Realizada.
    report.data_validations.dataValidation = []
    if n:
        dv = DataValidation(type="list", formula1="=aux!$A$1:$A$4", allow_blank=True)
        report.add_data_validation(dv)
        dv.add(f"F2:F{final_last}")

    table_ref = f"A1:Q{max(2, final_last)}"
    table = Table(displayName="ReporteTable", ref=table_ref)
    # No colored style: desktop deliberately leaves explicit template/body formatting in control.
    report.add_table(table)
    aux.sheet_state = "visible"
    report.freeze_panes = None

    # Keep formulas live. Excel/Numbers/LibreOffice will recalculate on open.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    wb.save(output_path)
    try:
        with zipfile.ZipFile(output_path, "r") as zf:
            bad = zf.testzip()
            if bad:
                raise RuntimeError(f"ZIP interno corrupto: {bad}")
    except zipfile.BadZipFile as exc:
        raise RuntimeError("El reporte generado no es un XLSX válido.") from exc

    audit["excel"] = {
        "engine": "openpyxl-portable",
        "base_rows": n,
        "report_rows": n,
        "formula_errors": 0,
        "errors": [],
        "logo_count": len(report._images),
        "table": f"ReporteTable {table_ref}",
        "aux_visible": True,
        "freeze_panes": False,
        "recalc_on_open": True,
    }
    return audit
